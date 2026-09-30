import pytest
from pathlib import Path
from app.services.ingestion.loaders.word_loader import WordLoader


def _has_docx_deps():
    try:
        import docx
        import docx2txt
        return True
    except ImportError:
        return False


class TestWordLoader:
    """Word加载器测试"""

    @pytest.mark.skipif(not _has_docx_deps(), reason="python-docx and docx2txt not installed")
    def test_load_word_document(self, tmp_path):
        """测试加载Word文档"""
        docx_file = tmp_path / "test.docx"
        from docx import Document
        doc = Document()
        doc.add_paragraph("Hello World")
        doc.add_paragraph("Test content")
        doc.core_properties.title = "Test Title"
        doc.core_properties.author = "Test Author"
        doc.save(str(docx_file))

        loader = WordLoader(str(docx_file))
        result = loader.load()

        assert "Hello World" in result.content
        assert "Test content" in result.content
        assert result.file_type == "docx"
        assert result.source == str(docx_file)
        assert result.metadata["title"] == "Test Title"
        assert result.metadata["author"] == "Test Author"

    @pytest.mark.skipif(not _has_docx_deps(), reason="python-docx and docx2txt not installed")
    def test_load_empty_word(self, tmp_path):
        """测试加载空Word文档"""
        docx_file = tmp_path / "empty.docx"
        from docx import Document
        doc = Document()
        doc.save(str(docx_file))

        loader = WordLoader(str(docx_file))
        result = loader.load()

        assert result.content.strip() == ""
        assert result.file_type == "docx"

    def test_load_nonexistent_file(self):
        """测试加载不存在的文件"""
        loader = WordLoader("/nonexistent/file.docx")
        with pytest.raises(FileNotFoundError):
            loader.load()

    @pytest.mark.skipif(not _has_docx_deps(), reason="python-docx and docx2txt not installed")
    def test_load_missing_docx2txt(self, tmp_path, monkeypatch):
        """测试缺少docx2txt依赖时的错误提示"""
        docx_file = tmp_path / "test.docx"
        from docx import Document
        doc = Document()
        doc.save(str(docx_file))

        import builtins
        import sys

        real_import = builtins.__import__

        def mock_import(name, *args, **kwargs):
            if name == "docx2txt":
                raise ImportError("mocked")
            return real_import(name, *args, **kwargs)

        monkeypatch.delitem(sys.modules, "docx2txt", raising=False)
        monkeypatch.setattr("builtins.__import__", mock_import)

        loader = WordLoader(str(docx_file))
        with pytest.raises(ImportError):
            loader.load()
