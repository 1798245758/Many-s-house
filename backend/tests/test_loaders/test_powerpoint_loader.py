import pytest
from pathlib import Path
from app.services.ingestion.loaders.powerpoint_loader import PowerPointLoader


class TestPowerPointLoader:
    """PowerPoint加载器测试"""

    def test_load_powerpoint_document(self, tmp_path):
        """测试加载PowerPoint文档"""
        pptx_file = tmp_path / "test.pptx"
        try:
            from pptx import Presentation
            prs = Presentation()
            slide_layout = prs.slide_layouts[1]  # Title and Content
            slide = prs.slides.add_slide(slide_layout)
            title = slide.shapes.title
            title.text = "Test Title"
            body = slide.placeholders[1]
            body.text = "Hello World\nTest content"
            prs.save(str(pptx_file))
        except ImportError:
            pytest.skip("python-pptx not installed")

        loader = PowerPointLoader(str(pptx_file))
        result = loader.load()

        assert "Test Title" in result.content
        assert "Hello World" in result.content
        assert "Test content" in result.content
        assert result.file_type == "pptx"
        assert result.source == str(pptx_file)
        assert "slides" in result.structure
        assert len(result.structure["slides"]) == 1

    def test_load_empty_powerpoint(self, tmp_path):
        """测试加载空PowerPoint文档"""
        pptx_file = tmp_path / "empty.pptx"
        try:
            from pptx import Presentation
            prs = Presentation()
            prs.save(str(pptx_file))
        except ImportError:
            pytest.skip("python-pptx not installed")

        loader = PowerPointLoader(str(pptx_file))
        result = loader.load()

        assert result.content.strip() == ""
        assert result.file_type == "pptx"
        assert "slides" in result.structure
        assert len(result.structure["slides"]) == 0

    def test_load_nonexistent_file(self):
        """测试加载不存在的文件"""
        with pytest.raises(FileNotFoundError):
            loader = PowerPointLoader("/nonexistent/file.pptx")
            loader.load()

    def test_load_multiple_slides(self, tmp_path):
        """测试加载包含多个幻灯片的PowerPoint文档"""
        pptx_file = tmp_path / "multi.pptx"
        try:
            from pptx import Presentation
            prs = Presentation()
            for i in range(3):
                slide_layout = prs.slide_layouts[1]
                slide = prs.slides.add_slide(slide_layout)
                title = slide.shapes.title
                title.text = f"Slide {i+1}"
                body = slide.placeholders[1]
                body.text = f"Content for slide {i+1}"
            prs.save(str(pptx_file))
        except ImportError:
            pytest.skip("python-pptx not installed")

        loader = PowerPointLoader(str(pptx_file))
        result = loader.load()

        assert "Slide 1" in result.content
        assert "Slide 3" in result.content
        assert len(result.structure["slides"]) == 3
        assert result.structure["slides"][0]["slide_number"] == 1
        assert result.structure["slides"][2]["slide_number"] == 3

    def test_load_with_table(self, tmp_path):
        """测试加载包含表格的PowerPoint文档"""
        pptx_file = tmp_path / "table.pptx"
        try:
            from pptx import Presentation
            from pptx.util import Inches
            prs = Presentation()
            slide_layout = prs.slide_layouts[1]
            slide = prs.slides.add_slide(slide_layout)
            rows, cols = 2, 3
            left = Inches(1.0)
            top = Inches(2.0)
            width = Inches(8.0)
            height = Inches(2.0)
            table_shape = slide.shapes.add_table(rows, cols, left, top, width, height)
            table = table_shape.table
            for i in range(cols):
                table.cell(0, i).text = f"Header {i+1}"
            for i in range(1, rows):
                for j in range(cols):
                    table.cell(i, j).text = f"Cell {i},{j}"
            prs.save(str(pptx_file))
        except ImportError:
            pytest.skip("python-pptx not installed")

        loader = PowerPointLoader(str(pptx_file))
        result = loader.load()

        assert "Header 1" in result.content
        assert "Cell 1,2" in result.content
        assert result.file_type == "pptx"

    def test_load_missing_pptx_dependency(self, tmp_path, monkeypatch):
        """测试缺少python-pptx依赖时的错误提示"""
        pptx_file = tmp_path / "test.pptx"
        try:
            from pptx import Presentation
            prs = Presentation()
            prs.save(str(pptx_file))
        except ImportError:
            pytest.skip("python-pptx not installed")

        import sys
        monkeypatch.setitem(sys.modules, "pptx", None)
        monkeypatch.setitem(sys.modules, "pptx.util", None)

        loader = PowerPointLoader(str(pptx_file))
        with pytest.raises(ImportError):
            loader.load()
