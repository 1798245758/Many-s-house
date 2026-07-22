from pathlib import Path
from typing import Any, Dict, List
from app.services.ingestion.loaders.base import BaseLoader, DocumentResult


class ExcelLoader(BaseLoader):
    """Excel文件加载器"""
    
    def load(self) -> DocumentResult:
        """加载Excel文档"""
        if not self.file_path.exists():
            raise FileNotFoundError(f"文件不存在: {self.file_path}")
        
        try:
            from openpyxl import load_workbook
        except ImportError:
            raise ImportError("请安装openpyxl: pip install openpyxl")
        
        wb = load_workbook(str(self.file_path), read_only=True, data_only=True)
        metadata = self._get_file_metadata()
        content_parts = []
        structure: Dict[str, Any] = {"sheets": []}
        
        for sheet_name in wb.sheetnames:
            ws = wb[sheet_name]
            sheet_data = {
                "name": sheet_name,
                "rows": ws.max_row or 0,
                "columns": ws.max_column or 0,
            }
            structure["sheets"].append(sheet_data)
            
            content_parts.append(f"[工作表: {sheet_name}]")
            for row in ws.iter_rows(values_only=True):
                row_data = [str(cell) if cell is not None else "" for cell in row]
                if any(row_data):
                    content_parts.append("\t".join(row_data))
        
        wb.close()
        
        content = "\n".join(content_parts)
        
        return DocumentResult(
            content=content,
            metadata=metadata,
            structure=structure,
            source=str(self.file_path),
            file_type="xlsx",
        )
