from pathlib import Path
from typing import Any, Dict, List
from app.services.ingestion.loaders.base import BaseLoader, DocumentResult


class PowerPointLoader(BaseLoader):
    """PowerPoint文件加载器"""

    def load(self) -> DocumentResult:
        """加载PowerPoint文档"""
        if not self.file_path.exists():
            raise FileNotFoundError(f"文件不存在: {self.file_path}")

        try:
            from pptx import Presentation
        except ImportError:
            raise ImportError("请安装python-pptx: pip install python-pptx")

        prs = Presentation(str(self.file_path))
        metadata = self._get_file_metadata()
        content_parts: List[str] = []
        structure: Dict[str, Any] = {"slides": []}

        for i, slide in enumerate(prs.slides, 1):
            slide_content: Dict[str, Any] = {
                "slide_number": i,
                "texts": [],
            }

            content_parts.append(f"[幻灯片 {i}]")

            for shape in slide.shapes:
                if hasattr(shape, "text") and shape.text:
                    text = shape.text.strip()
                    if text:
                        content_parts.append(text)
                        slide_content["texts"].append(text)

                if shape.has_table:
                    table = shape.table
                    for row in table.rows:
                        row_data = [cell.text.strip() for cell in row.cells]
                        if any(row_data):
                            content_parts.append("\t".join(row_data))

            structure["slides"].append(slide_content)

        content = "\n".join(content_parts)

        return DocumentResult(
            content=content,
            metadata=metadata,
            structure=structure,
            source=str(self.file_path),
            file_type="pptx",
        )
